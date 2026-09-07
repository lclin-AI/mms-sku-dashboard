"""
Verify one SKU/date on the dashboard against the LIVE authoritative sources.

For a (sku, delivery-date) it prints, side by side:
  metric        dashboard (Supabase)     live source                 match?
  sold_standard  mms_sku_daily            MMS Daily Order Report
  sold_express   mms_express_daily        MMS consignments (prm=C) + details
  PO / GR        imax_daily               iMAX getPurchaseOrderInfo (bookedDate)
  disposal       imax_daily               iMAX process log (R-loc net, 4PM shift)
  IIMS           imax_iims_sku            IIMS stock-levels API

Run locally (needs MMS/iMAX/IIMS which are internal). Remaining (4PM snapshot)
is a point-in-time capture that cannot be re-fetched historically, so it is not
re-verified here.

  python verify.py --sku 00001 --date 2026-09-04
"""
import argparse, io, json, os, sys
from collections import defaultdict
from datetime import datetime, timezone, timedelta

import mms_login, requests
from openpyxl import load_workbook

HK = timezone(timedelta(hours=8))
HERE = os.path.dirname(os.path.abspath(__file__))
MMS = "https://merchant-web.shoalter.com"
IMAX = "https://imax.hktvmall.com/hktv_imax"
IMAX_CFG = r"C:\Users\lclin\imax-stock-slack\config.json"
ACTIVE = ["CONFIRMED", "ACKNOWLEDGED", "PACKED", "PICKED", "PICKEDUP_FROM_MERCHANT",
          "IN_HUB", "DISPATCHED", "IN_STORE", "IN_LOCKER", "MERCHANT_SHIPPED",
          "FAIL_TO_DELIVER", "HOLD_BY_CS", "RELEASE_BY_CS", "RECEIVED_BY_CUSTOMER",
          "ORDER_COMPLETE"]
DISPOSAL_LOC = None  # loaded from the disposal module


def load_run_env():
    p = os.path.join(HERE, "run.env")
    if os.path.exists(p):
        for line in open(p, encoding="utf-8"):
            s = line.strip()
            if s and not s.startswith("#") and "=" in s:
                k, v = s.split("=", 1)
                os.environ[k.strip()] = v.strip()


def anon_headers():
    cfg = open(os.path.join(HERE, "config.js"), encoding="utf-8").read()
    import re
    key = re.search(r'SUPABASE_ANON_KEY:\s*"([^"]+)"', cfg).group(1)
    url = re.search(r'SUPABASE_URL:\s*"([^"]+)"', cfg).group(1)
    return url, {"apikey": key, "Authorization": "Bearer " + key}


def imax_get(path, params):
    s = requests.Session()
    s.headers.update({"User-Agent": "Mozilla/5.0", "Accept": "application/json", "Referer": IMAX + "/"})
    sf = os.path.join(HERE, "imax_session.json")
    if os.path.exists(sf):
        s.cookies.update(json.load(open(sf)))
        r = s.get(IMAX + path, params=params, timeout=120)
        if r.status_code == 200:
            return r
    d = json.load(open(IMAX_CFG, encoding="utf-8"))
    s.cookies.clear()
    s.post(IMAX + "/login", data={"username": d["imaxUser"], "password": d["imaxPass"]},
           headers={"Content-Type": "application/x-www-form-urlencoded"}, timeout=30)
    json.dump({c.name: c.value for c in s.cookies}, open(sf, "w"))
    return s.get(IMAX + path, params=params, timeout=120)


def biz_day(ms):
    dt = datetime.fromtimestamp(ms / 1000, HK)
    if dt.hour >= 16:
        dt = dt + timedelta(days=1)
    return dt.strftime("%Y-%m-%d")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sku", required=True)
    ap.add_argument("--date", required=True, help="delivery date YYYY-MM-DD")
    ap.add_argument("--store", default="B0812001")
    a = ap.parse_args()
    load_run_env()
    store, sku, date = a.store, a.sku, a.date
    full = f"{store}_S_{sku}"
    ymd = date.replace("-", "")

    sb_url, H = anon_headers()

    def sb(table, q):
        return requests.get(f"{sb_url}/rest/v1/{table}?{q}", headers=H, timeout=60).json()

    # ---- DASHBOARD (Supabase) ----
    d_std = sum(r["qty"] for r in sb("mms_sku_daily",
               f"store_code=eq.{store}&sku_id=eq.{sku}&delivery_date=eq.{date}&select=qty"))
    d_exp = sum(r["qty"] for r in sb("mms_express_daily",
               f"store_code=eq.{store}&sku_id=eq.{sku}&delivery_date=eq.{date}&select=qty"))
    id_rows = sb("imax_daily", f"store_code=eq.{store}&sku_id=eq.{sku}&date=eq.{date}&select=po_qty,gr_qty,disposal_qty")
    d_po = id_rows[0]["po_qty"] if id_rows else 0
    d_gr = id_rows[0]["gr_qty"] if id_rows else 0
    d_disp = (id_rows[0]["disposal_qty"] if id_rows else None)
    iims_rows = sb("imax_iims_sku", f"store_code=eq.{store}&sku_id=eq.{sku}&select=quantity")
    d_iims = iims_rows[0]["quantity"] if iims_rows else None

    # ---- LIVE ----
    tok = mms_login.get_token()
    m = requests.Session(); m.headers.update({"Authorization": "Bearer " + tok, "Content-Type": "application/json"})

    # standard sold: Daily Order Report
    meta = m.get(f"https://merchant-web.shoalter.com/order/HKTV/{store}/DAILY/report",
                 params={"dates": ",".join((datetime.strptime(date, "%Y-%m-%d").date() - timedelta(days=i)).strftime("%Y%m%d") for i in range(0, 12))}).json()["data"]["data"]
    l_std = 0.0
    for fn in meta:
        b = m.get(f"https://merchant-web.shoalter.com/order/HKTV/DAILY/{fn}/downloadreport").content
        wb = load_workbook(io.BytesIO(b), data_only=True); ws = wb.active; hdr = None
        for row in ws.iter_rows(values_only=True):
            if hdr is None:
                if any(str(c).strip() == "SKU ID" for c in row if c is not None):
                    hdr = [str(c).replace("\n", " ").strip() if c else "" for c in row]
                continue
            r = dict(zip(hdr, row))
            if str(r.get("SKU ID") or "") == sku and str(r.get("Delivery Date") or "")[:10] == date \
               and "CANCEL" not in str(r.get("Status") or "").upper():
                l_std += float(r.get("Qty (Q)") or 0)
        wb.close()

    # express sold: consignments (prm C) delivering date, details
    def dms(y_m_d, end=False):
        dt = datetime.strptime(y_m_d, "%Y-%m-%d").replace(tzinfo=HK)
        if end: dt = dt.replace(hour=23, minute=59, second=59)
        return int(dt.timestamp() * 1000)
    wh = [f"{store}{i:02d}" for i in range(1, 100)]
    codes = []; p = 1
    while True:
        body = {"storefrontStoreCodes": [store], "productReadyMethods": ["C"], "warehouseCodes": wh,
                "startDate": dms(date), "endDate": dms(date, True), "status": ACTIVE,
                "searchDateType": "DELIVERY_DATE", "deliveryMode": "STANDARD_DELIVERY",
                "sortColumn": "ISSUE_DATE", "sortDirection": "DESC", "searchType": "ORDER_ID",
                "searchKeyword": "", "pageNumber": p, "pageSize": 1000}
        resp = m.post(f"{MMS}/order/v2/consignments", data=json.dumps(body), timeout=90).json().get("response") or {}
        codes += [x.get("consignmentCode") for x in resp.get("data") or []]
        if p >= (resp.get("pagination") or {}).get("numberOfPages", 1): break
        p += 1
    from concurrent.futures import ThreadPoolExecutor
    l_exp = 0.0
    def det(c):
        try:
            e = (m.get(f"{MMS}/order/v2/{c}/consignmentDetails", timeout=60).json().get("data") or {}).get("consignmentEntries") or []
            return sum(float(x.get("quantity") or 0) for x in e if str(x.get("skuId", "")).split("_S_")[-1] == sku)
        except Exception:
            return 0.0
    with ThreadPoolExecutor(max_workers=24) as ex:
        l_exp = sum(ex.map(det, codes))

    # PO/GR: iMAX purchase orders
    po = imax_get("/frontend/getPurchaseOrderInfo", {"rmCode": "E0059"}).json()
    l_po = l_gr = 0.0
    for x in po:
        if str(x.get("materialCode")) == f"{store}-{sku}" and str(x.get("status")) != "CANCELLED":
            bd = x.get("bookedDate")
            if bd and datetime.fromtimestamp(bd / 1000, HK).strftime("%Y-%m-%d") == date:
                l_po += float(x.get("orderQuantity") or 0)
                l_gr += float(x.get("receivedQuantity") or 0)

    # disposal: process log for this SKU
    import imax_disposal_to_supabase as dm
    now = datetime.now(HK)
    frm = int((now - timedelta(days=14)).replace(hour=0, minute=0, second=0, microsecond=0).timestamp() * 1000)
    log = imax_get("/frontend/getMaterialProcessLog", {"from": frm, "to": int(now.timestamp() * 1000), "material": full.replace("_S_", "-").rsplit("-", 1)[0] + "-" + sku}).json()
    l_disp = 0.0
    for x in (log if isinstance(log, list) else log.get("data", [])):
        loc = str(x.get("location") or "").upper(); tx = x.get("txDescription")
        if loc in dm.DISPOSAL_LOCATIONS and tx != "sap_disposal" and biz_day(x.get("processDate")) == date:
            l_disp += float(x.get("qty") or 0)

    # IIMS live
    ii = requests.get(f"https://iims-restful.shoalter.com/iims/s2s/v2/hybris/products/{full}/stock-levels", timeout=20).json().get("data", {})
    l_iims = ii.get("quantity")

    def line(name, dash, live):
        ok = "OK" if (dash is not None and live is not None and abs(float(dash) - float(live)) < 0.5) else "DIFF"
        print(f"  {name:16} dashboard={str(dash):>10}   live={str(live):>10}   {ok}")

    print(f"\n=== VERIFY {sku} delivery {date} ({store}) ===")
    line("sold_standard", round(d_std), round(l_std))
    line("sold_express", round(d_exp), round(l_exp))
    line("sold_total", round(d_std + d_exp), round(l_std + l_exp))
    line("PO", d_po, round(l_po))
    line("GR", d_gr, round(l_gr))
    line("disposal", d_disp, round(l_disp))
    line("IIMS", d_iims, l_iims)


if __name__ == "__main__":
    main()
